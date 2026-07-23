"""
خروجی اکسل گزارش تفصیلی ماهانه - نسخه ۲
با ستون‌های متعدد ورود/خروج
"""
from datetime import date
from typing import Dict, List
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import jdatetime


class DetailedExcelExporterV2:
    """خروجی اکسل گزارش تفصیلی - نسخه ۲"""

    def __init__(self, output_dir: str = "exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)

    def export_detailed_report(self, report: Dict) -> str:
        """خروجی گزارش تفصیلی یک کارمند به اکسل"""
        wb = Workbook()

        ws_daily = wb.active
        ws_daily.title = "گزارش روزانه"
        self._create_daily_sheet(ws_daily, report)

        ws_summary = wb.create_sheet("خلاصه ماهانه")
        self._create_summary_sheet(ws_summary, report)

        emp = report['employee']
        month_name = report['month_name']
        year = report['year']
        filename = self.output_dir / f"گزارش_تفصیلی_V2_{emp['full_name']}_{month_name}_{year}.xlsx"
        wb.save(filename)
        return str(filename)

    def export_all_employees_report(self, reports: List[Dict], year: int, month: int, month_name: str) -> str:
        """خروجی گزارش همه کارمندان در یک فایل اکسل"""
        wb = Workbook()

        ws_summary = wb.active
        ws_summary.title = "خلاصه کلی"
        self._create_all_summary_sheet(ws_summary, reports, year, month, month_name)

        for report in reports:
            emp = report['employee']
            sheet_name = f"{emp['full_name'][:25]} ({emp['user_id']})"
            ws = wb.create_sheet(sheet_name[:31])
            self._create_daily_sheet(ws, report)

            ws_sum = wb.create_sheet(f"خلاصه {emp['user_id']}")
            self._create_summary_sheet(ws_sum, report)

        filename = self.output_dir / f"گزارش_کلی_V2_{month_name}_{year}.xlsx"
        wb.save(filename)
        return str(filename)

    def _create_daily_sheet(self, ws, report: Dict):
        """ایجاد شیت گزارش روزانه با ستون‌های متعدد ورود/خروج"""
        emp = report['employee']

        ws.append([f"گزارش تفصیلی ماهانه V2 - {emp['full_name']} ({emp['user_id']})"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=15)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([f"ماه: {report['month_name']} {report['year']}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=15)
        ws['A2'].alignment = Alignment(horizontal='center')

        ws.append([])

        # ✅ هدر با ستون‌های متعدد ورود/خروج
        headers = [
            'تاریخ', 'روز', 'وضعیت روز', 'وضعیت فرد',
            'ورود ۱', 'خروج ۱', 'ورود ۲', 'خروج ۲', 'ورود ۳', 'خروج ۳',
            'وضعیت تردد', 'کارکرد', 'اضافی', 'کسری'
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for day in report['days']:
            def fmt_time(dt):
                return dt.strftime('%H:%M') if dt else '-'

            def fmt_hours(h):
                if h == 0:
                    return '-'
                hours = int(h)
                minutes = int((h - hours) * 60)
                return f"{hours:02d}:{minutes:02d}"

            # ✅ استخراج جفت‌های ورود/خروج
            pairs = day.get('attendance_pairs', [])
            enter1 = fmt_time(pairs[0]['enter']) if len(pairs) > 0 else '-'
            exit1 = fmt_time(pairs[0]['exit']) if len(pairs) > 0 else '-'
            enter2 = fmt_time(pairs[1]['enter']) if len(pairs) > 1 else '-'
            exit2 = fmt_time(pairs[1]['exit']) if len(pairs) > 1 else '-'
            enter3 = fmt_time(pairs[2]['enter']) if len(pairs) > 2 else '-'
            exit3 = fmt_time(pairs[2]['exit']) if len(pairs) > 2 else '-'

            attendance_str = day['attendance_status']
            if attendance_str == 'بدون تردد':
                attendance_str = '-'

            ws.append([
                day['jalali_date'],
                day['day_name'],
                day['day_status'],
                day['person_status_name'],
                enter1, exit1,
                enter2, exit2,
                enter3, exit3,
                attendance_str,
                fmt_hours(day['work_hours']),
                fmt_hours(day['surplus']),
                fmt_hours(day['deficit'])
            ])

        self._auto_adjust_column_width(ws)

    def _create_summary_sheet(self, ws, report: Dict):
        """ایجاد شیت خلاصه ماهانه"""
        emp = report['employee']
        summary = report['summary']

        ws.append([f"خلاصه ماهانه - {emp['full_name']} ({emp['user_id']})"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([f"ماه: {report['month_name']} {report['year']}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
        ws['A2'].alignment = Alignment(horizontal='center')

        ws.append([])

        ws.append(['موظفی'])
        ws['A4'].font = Font(bold=True, size=12)
        ws.append(['روزهای موظفی', summary['duty_days']])
        ws.append(['ساعات موظفی', self._fmt_hours(summary['duty_hours'])])

        ws.append([])

        ws.append(['وضعیت روزها'])
        ws['A7'].font = Font(bold=True, size=12)
        ws.append(['حضور', summary['present_days']])
        ws.append(['مرخصی', summary['leave_days']])
        ws.append(['غیبت', summary['absent_days']])
        ws.append(['استراحت', summary['rest_days']])
        ws.append(['تعطیل', summary['holiday_days']])
        ws.append(['جمعه کاری', summary['friday_work_days']])

        ws.append([])

        ws.append(['ساعات کاری'])
        ws['A15'].font = Font(bold=True, size=12)
        ws.append(['کارکرد ماهانه', self._fmt_hours(summary['total_work_hours'])])
        ws.append(['ساعات صبح', self._fmt_hours(summary['total_morning'])])
        ws.append(['ساعات عصر', self._fmt_hours(summary['total_evening'])])
        ws.append(['ساعات شب', self._fmt_hours(summary['total_night'])])

        ws.append([])

        ws.append(['اضافه کاری و کسری'])
        ws['A21'].font = Font(bold=True, size=12)
        ws.append(['اضافه کاری روزانه (مبنای 7:20)', self._fmt_hours(summary['total_surplus'])])
        ws.append(['کسری کار (مبنای 7:20)', self._fmt_hours(summary['total_deficit'])])
        ws.append(['اضافه کاری هفتگی', self._fmt_hours(summary['weekly_overtime'])])
        ws.append(['جمعه کاری', self._fmt_hours(summary['friday_work_hours'])])

        self._auto_adjust_column_width(ws)

    def _create_all_summary_sheet(self, ws, reports: List[Dict], year: int, month: int, month_name: str):
        """ایجاد شیت خلاصه کلی همه کارمندان"""
        ws.append([f"خلاصه گزارش ماهانه V2 - {month_name} {year}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([f"تعداد کارمندان: {len(reports)}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
        ws['A2'].alignment = Alignment(horizontal='center')

        ws.append([])

        headers = [
            'ردیف', 'کد', 'نام کامل', 'دپارتمان',
            'روز موظفی', 'ساعت موظفی', 'حضور', 'مرخصی', 'غیبت',
            'کارکرد', 'اضافی', 'کسری'
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for i, report in enumerate(reports, 1):
            emp = report['employee']
            summary = report['summary']

            ws.append([
                i,
                emp['user_id'],
                emp['full_name'],
                emp['department'],
                summary['duty_days'],
                self._fmt_hours(summary['duty_hours']),
                summary['present_days'],
                summary['leave_days'],
                summary['absent_days'],
                self._fmt_hours(summary['total_work_hours']),
                self._fmt_hours(summary['total_surplus']),
                self._fmt_hours(summary['total_deficit'])
            ])

        self._auto_adjust_column_width(ws)

    def _fmt_hours(self, h: float) -> str:
        """فرمت ساعات"""
        if h == 0:
            return '00:00'
        hours = int(h)
        minutes = int((h - hours) * 60)
        return f"{hours:02d}:{minutes:02d}"

    def _auto_adjust_column_width(self, ws, max_width: int = 25):
        """تنظیم خودکار عرض ستون‌ها"""
        for col_idx in range(1, ws.max_column + 1):
            column_letter = get_column_letter(col_idx)
            max_length = 0

            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    try:
                        cell_len = len(str(cell.value))
                        if cell_len > max_length:
                            max_length = cell_len
                    except:
                        pass

            adjusted_width = min(max_length + 2, max_width)
            ws.column_dimensions[column_letter].width = max(adjusted_width, 12)