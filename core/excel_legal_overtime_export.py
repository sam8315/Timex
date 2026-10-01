"""خروجی اکسل گزارش اضافه‌کار قانونی."""
from typing import Dict

from openpyxl.styles import Font, PatternFill, Alignment

from core.excel_detailed_export_v2 import DetailedExcelExporterV2


class LegalOvertimeExcelExporter(DetailedExcelExporterV2):
    """مثل V2 با ستون هفتگی به‌جای کسری روزانه و خلاصهٔ پرداختی."""

    def _create_daily_sheet(self, ws, report: Dict):
        emp = report['employee']
        headers = [
            'تاریخ', 'روز', 'وضعیت روز', 'وضعیت فرد',
            'ورود ۱', 'خروج ۱', 'ورود ۲', 'خروج ۲', 'ورود ۳', 'خروج ۳',
            'وضعیت تردد', 'کارکرد', 'هفتگی',
        ]
        last_col = len(headers)

        ws.append([
            f"گزارش اضافه‌کار قانونی - {emp['full_name']} ({emp['user_id']})"
        ])
        ws.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=last_col)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([f"ماه: {report['month_name']} {report['year']}"])
        ws.merge_cells(
            start_row=2, start_column=1, end_row=2, end_column=last_col)
        ws['A2'].alignment = Alignment(horizontal='center')
        ws.append([])
        ws.append(headers)

        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        data_start_row = 5  # ردیف اول داده بعد از عنوان/هدر
        week_col = 13  # ستون «هفتگی»
        merge_ranges = []

        for idx, day in enumerate(report['days']):
            def fmt_time(dt):
                return dt.strftime('%H:%M') if dt else '-'

            def fmt_hours(h):
                if not h:
                    return '-'
                return self._fmt_hours(h)

            pairs = day.get('attendance_pairs', [])
            enter1 = fmt_time(pairs[0]['enter']) if len(pairs) > 0 else '-'
            exit1 = fmt_time(pairs[0]['exit']) if len(pairs) > 0 else '-'
            enter2 = fmt_time(pairs[1]['enter']) if len(pairs) > 1 else '-'
            exit2 = fmt_time(pairs[1]['exit']) if len(pairs) > 1 else '-'
            enter3 = fmt_time(pairs[2]['enter']) if len(pairs) > 2 else '-'
            exit3 = fmt_time(pairs[2]['exit']) if len(pairs) > 2 else '-'

            attendance_str = day.get('attendance_status') or '-'
            if attendance_str == 'بدون تردد':
                attendance_str = '-'

            week_text = ''
            if day.get('week_rowspan') and day.get('week_block'):
                wb = day['week_block']
                week_text = (
                    f"کارکرد: {fmt_hours(wb.get('work_hours'))}\n"
                    f"موظفی: {fmt_hours(wb.get('duty_hours'))} ({wb.get('duty_note')})\n"
                    f"اضافه‌کار هفتگی: {fmt_hours(wb.get('weekly_overtime'))}\n"
                    f"{wb.get('formula') or ''}"
                )
                start_row = data_start_row + idx
                end_row = start_row + int(day['week_rowspan']) - 1
                if end_row > start_row:
                    merge_ranges.append((start_row, end_row))

            ws.append([
                day['jalali_date'],
                day['day_name'],
                day['day_status'],
                day['person_status_name'],
                enter1, exit1,
                enter2, exit2,
                enter3, exit3,
                attendance_str,
                fmt_hours(day.get('work_hours')),
                week_text,
            ])

        for start_row, end_row in merge_ranges:
            ws.merge_cells(
                start_row=start_row, start_column=week_col,
                end_row=end_row, end_column=week_col,
            )
            cell = ws.cell(row=start_row, column=week_col)
            cell.alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=True)

        self._auto_adjust_column_width(ws)
        ws.column_dimensions['M'].width = 28

    def _create_summary_sheet(self, ws, report: Dict):
        emp = report['employee']
        summary = report['summary']

        ws.append([
            f"خلاصه اضافه‌کار قانونی - {emp['full_name']} ({emp['user_id']})"
        ])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([f"ماه: {report['month_name']} {report['year']}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
        ws['A2'].alignment = Alignment(horizontal='center')
        ws.append([])

        def section(title):
            ws.append([title])
            row = ws.max_row
            ws.cell(row=row, column=1).font = Font(bold=True, size=12)
            ws.append([])

        def entry(label, value):
            ws.append([label, value])

        section('کل کارکرد و موظفی')
        entry('کل کارکرد', self._fmt_hours(summary['total_work_hours']))
        entry('تعداد روز کارکرد', summary.get('total_work_days', 0))
        entry('روزهای موظفی', summary['duty_days'])
        entry('ساعات موظفی', self._fmt_hours(summary['duty_hours']))

        section('اضافه‌کار (کارکرد − جمعه − تعطیل‌کاری − Σ موظفی)')
        entry('اضافه‌کار هفتگی', self._fmt_hours(summary['weekly_overtime']))
        entry('جمع اضافه‌کار', self._fmt_hours(summary['overtime_total']))

        section('تعطیل‌کاری (از مبنای هفتگی کسر)')
        entry('ساعت تعطیل‌کاری', self._fmt_hours(summary.get('holiday_work_hours') or 0))
        entry('تعداد روز تعطیل‌کاری', summary.get('holiday_work_days', 0))

        section('مجموع اضافه‌کار و تعطیل‌کاری')
        entry('مجموع', self._fmt_hours(summary.get('overtime_holiday_total') or 0))

        section('جمعه‌کاری (از مبنای هفتگی کسر)')
        entry('ساعت جمعه‌کاری', self._fmt_hours(summary['friday_work_hours']))
        entry('تعداد روز جمعه‌کاری', summary['friday_work_days'])

        section('تفکیک صبح / عصر / شب')
        entry('ساعات صبح', self._fmt_hours(summary.get('total_morning') or 0))
        entry('درصد صبح', f"{summary.get('morning_percent', 0)}٪")
        entry('ساعات عصر', self._fmt_hours(summary.get('total_evening') or 0))
        entry('درصد عصر', f"{summary.get('evening_percent', 0)}٪")
        entry('ساعات شب', self._fmt_hours(summary.get('total_night') or 0))
        entry('درصد شب', f"{summary.get('night_percent', 0)}٪")

        section('وضعیت روزها و مرخصی/مأموریت')
        entry('حضور', summary.get('present_days', 0))
        entry('مرخصی روزانه', summary.get('leave_days', 0))
        entry('مرخصی ساعتی', self._fmt_hours(summary.get('hourly_leave_hours') or 0))
        entry('مأموریت روزانه', summary.get('mission_days', 0))
        entry('مأموریت ساعتی', self._fmt_hours(summary.get('hourly_mission_hours') or 0))
        entry('عدم حضور', summary.get('absent_days', 0))
        entry('استراحت', summary.get('rest_days', 0))
        entry('تعطیل', summary.get('holiday_days', 0))
        entry('کسر کار ماهانه', self._fmt_hours(summary.get('monthly_deficit') or 0))

        self._auto_adjust_column_width(ws)
