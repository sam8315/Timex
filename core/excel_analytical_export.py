"""
ماژول خروجی اکسل گزارش تحلیلی
"""
from datetime import date
from typing import Dict
from pathlib import Path
import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class AnalyticalExcelExporter:
    """خروجی اکسل گزارش تحلیلی"""

    def __init__(self, output_dir: str = "exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)

    def export_analytical_report(self, report: Dict) -> str:
        """خروجی گزارش تحلیلی به اکسل"""
        wb = Workbook()
        wb.remove(wb.active)  # حذف شیت پیش‌فرض

        month_name = report['month_name']
        year = report['year']

        # شیت خلاصه
        ws_summary = wb.create_sheet("خلاصه")
        self._create_summary_sheet(ws_summary, report)

        # شیت هر گروه
        for group_name, users in report['groups'].items():
            ws = wb.create_sheet(group_name)
            self._create_group_sheet(ws, group_name, users, report)

        # ذخیره
        filename = self.output_dir / f"گزارش_تحلیلی_{month_name}_{year}.xlsx"
        wb.save(filename)
        return str(filename)

    def _create_summary_sheet(self, ws, report: Dict):
        """ایجاد شیت خلاصه"""
        ws.append([f"گزارش تحلیلی ماهانه - {report['month_name']} {report['year']}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([])
        ws.append(["گروه", "تعداد کاربران", "میانگین ساعات کاری", "میانگین کسری/اضافی"])

        for group_name, users in report['groups'].items():
            avg_work = sum(u['total_hours'] for u in users) / len(users) if users else 0
            avg_diff = sum(u['difference'] for u in users) / len(users) if users else 0

            ws.append([
                group_name,
                len(users),
                f"{int(avg_work)}:{int((avg_work - int(avg_work)) * 60):02d}",
                f"{int(avg_diff)}:{int((avg_diff - int(avg_diff)) * 60):02d}"
            ])

    def _create_group_sheet(self, ws, group_name: str, users: list, report: Dict):
        """ایجاد شیت یک گروه"""
        ws.append([f"گروه: {group_name}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
        ws['A1'].font = Font(bold=True, size=12)

        ws.append([])
        headers = [
            'ردیف', 'کد پرسنلی', 'نام کامل', 'حاضر', 'غایب', 'مرخصی',
            'روز موظفی', 'ساعت موظفی', 'صبح', 'عصر', 'شب', 'کل', 'کسری/اضافی'
        ]
        ws.append(headers)

        # استایل هدر
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # داده‌ها
        for i, u in enumerate(users, 1):
            def fmt_hours(h):
                if h == 0:
                    return "00:00"
                hours = int(h)
                minutes = int((h - hours) * 60)
                return f"{hours:02d}:{minutes:02d}"

            diff = u['difference']
            diff_str = f"+{fmt_hours(diff)}" if diff > 0 else fmt_hours(diff)

            ws.append([
                i,
                u['user_id'],
                u['full_name'],
                u['present_days'],
                u['absent_days'],
                u['leave_days'],
                u['required_days'],
                fmt_hours(u['required_hours']),
                fmt_hours(u['morning_hours']),
                fmt_hours(u['evening_hours']),
                fmt_hours(u['night_hours']),
                fmt_hours(u['total_hours']),
                diff_str
            ])

        # تنظیم عرض ستون‌ها
        for column in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in column)
            ws.column_dimensions[column[0].column_letter].width = max(max_length + 2, 12)