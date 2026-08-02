"""
ماژول خروجی اکسل گزارش‌ها
"""
from datetime import date
from typing import List, Dict, Optional
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import jdatetime


class ExcelExporter:
    """خروجی اکسل گزارش‌ها"""

    def __init__(self, output_dir: str = "exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)

    def _get_jalali_month_name(self, month: int) -> str:
        """دریافت نام ماه شمسی"""
        names = {
            1: 'فروردین', 2: 'اردیبهشت', 3: 'خرداد',
            4: 'تیر', 5: 'مرداد', 6: 'شهریور',
            7: 'مهر', 8: 'آبان', 9: 'آذر',
            10: 'دی', 11: 'بهمن', 12: 'اسفند'
        }
        return names.get(month, '')

    def _auto_adjust_column_width(self, ws, max_width: int = 25):
        """تنظیم خودکار عرض ستون‌ها"""
        for col_idx in range(1, ws.max_column + 1):
            column_letter = get_column_letter(col_idx)
            max_length = 0

            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    try:
                        cell_len = len(str(cell.value))
                        if cell_len > max_length:
                            max_length = cell_len
                    except:
                        pass

            adjusted_width = min(max_length + 2, max_width)
            ws.column_dimensions[column_letter].width = max(adjusted_width, 10)

    def export_monthly_report(self, reports: List[Dict], year: int, month: int) -> str:
        """خروجی گزارش ماهانه به اکسل"""
        wb = Workbook()
        ws = wb.active
        ws.title = "گزارش ماهانه"

        # هدر
        ws.append([f"گزارش ماهانه - {self._get_jalali_month_name(month)} {year}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=15)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([])

        # هدر ستون‌ها
        headers = [
            'ردیف', 'کد پرسنلی', 'نام کامل', 'دپارتمان', 'حاضر', 'غایب', 'تعطیل',
            'استحقاقی', 'استعلاجی', 'تشویقی', 'بدون حقوق', 'ماموریت', 'حضور کم', 'ویژه', 'ساعت کاری'
        ]
        ws.append(headers)

        # استایل هدر
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # داده‌ها
        for i, r in enumerate(reports, 1):
            work_str = f"{int(r['work_hours'])}:{int((r['work_hours'] % 1) * 60):02d}"

            ws.append([
                i,
                r['user_id'],
                r['full_name'],
                r.get('department', '-'),
                r['present_days'],
                r['absent_days'],
                r['holiday_days'],
                r['annual_leave'],
                r['sick_leave'],
                r['reward_leave'],
                r['unpaid_leave'],
                r['mission_days'],
                r['late_days'],
                r['special_days'],
                work_str
            ])

        # ✅ استفاده از متد جدید
        self._auto_adjust_column_width(ws, max_width=25)

        # ذخیره
        filename = self.output_dir / f"گزارش_ماهانه_{self._get_jalali_month_name(month)}_{year}.xlsx"
        wb.save(filename)
        return str(filename)

    def export_absent_report(self, reports: List[Dict], from_date: date, to_date: date) -> str:
        """خروجی گزارش غیبت‌ها به اکسل"""
        wb = Workbook()
        ws = wb.active
        ws.title = "گزارش غیبت‌ها"

        # هدر
        j_from = jdatetime.date.fromgregorian(date=from_date)
        j_to = jdatetime.date.fromgregorian(date=to_date)
        ws.append([f"گزارش غیبت‌ها - {j_from.strftime('%Y/%m/%d')} تا {j_to.strftime('%Y/%m/%d')}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([])

        # هدر ستون‌ها
        headers = ['ردیف', 'کد پرسنلی', 'نام کامل', 'دپارتمان', 'تعداد غیبت', 'تاریخ‌های غیبت']
        ws.append(headers)

        # استایل هدر
        header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # داده‌ها
        for i, r in enumerate(reports, 1):
            dates_str = ', '.join([
                jdatetime.date.fromgregorian(date=d).strftime('%Y/%m/%d')
                for d in r['absent_dates']
            ])

            ws.append([
                i,
                r['user_id'],
                r['full_name'],
                r.get('department', '-'),
                r['absent_count'],
                dates_str
            ])

        # ✅ استفاده از متد جدید
        self._auto_adjust_column_width(ws, max_width=40)

        # ذخیره
        filename = self.output_dir / f"گزارش_غیبت‌ها_{j_from.strftime('%Y%m%d')}_تا_{j_to.strftime('%Y%m%d')}.xlsx"
        wb.save(filename)
        return str(filename)

    def export_leave_report(self, reports: List[Dict], year: int, month: Optional[int] = None) -> str:
        """خروجی گزارش مرخصی‌ها به اکسل"""
        wb = Workbook()
        ws = wb.active
        ws.title = "گزارش مرخصی‌ها"

        # هدر
        period = f"{self._get_jalali_month_name(month)} {year}" if month else f"سال {year}"
        ws.append([f"گزارش مرخصی‌ها - {period}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([])

        # هدر ستون‌ها
        headers = ['ردیف', 'کد پرسنلی', 'نام کامل', 'دپارتمان', 'استحقاقی', 'استعلاجی', 'تشویقی', 'بدون حقوق', 'مجموع', 'تعداد درخواست']
        ws.append(headers)

        # استایل هدر
        header_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # داده‌ها
        for i, r in enumerate(reports, 1):
            ws.append([
                i,
                r['user_id'],
                r['full_name'],
                r.get('department', '-'),
                r['annual_leave'],
                r['sick_leave'],
                r['reward_leave'],
                r['unpaid_leave'],
                r['total_days'],
                r['total_requests']
            ])

        # ✅ استفاده از متد جدید
        self._auto_adjust_column_width(ws, max_width=20)

        # ذخیره
        filename = self.output_dir / f"گزارش_مرخصی‌ها_{year}_{month or 'کل'}.xlsx"
        wb.save(filename)
        return str(filename)