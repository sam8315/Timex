"""خروجی اکسل گزارش خام تردد."""
from io import BytesIO
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.raw_report import fmt_time


FONT_NAME = 'B Nazanin'
HEADER_FONT = Font(bold=True, size=10, name=FONT_NAME, color='000000')
DATA_FONT = Font(size=10, name=FONT_NAME)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)
THIN_SIDE = Side(style='thin', color='999999')
THIN = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)
HEADER_FILL = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
EVEN_FILL = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
OFF_FILL = PatternFill(start_color='E6E6E6', end_color='E6E6E6', fill_type='solid')


def _safe_sheet_name(value: str, used: set) -> str:
    cleaned = ''.join(char for char in str(value) if char not in '[]:*?/\\')
    base = (cleaned or 'گزارش')[:31]
    name = base
    suffix = 2
    while name in used:
        ending = f'-{suffix}'
        name = f'{base[:31 - len(ending)]}{ending}'
        suffix += 1
    used.add(name)
    return name


def _write_title(ws, title: str, month_name: str, year: int, last_col: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14, name=FONT_NAME, color='000000')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24
    if month_name and year:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        cell = ws.cell(row=2, column=1, value=f'{month_name} {year}')
        cell.font = Font(size=12, name=FONT_NAME)
        cell.alignment = Alignment(horizontal='center', vertical='center')


def _write_employee_header(ws, emp: dict):
    row = 4
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    cell = ws.cell(
        row=row,
        column=1,
        value=(
            f'نام و نام خانوادگی: {emp["full_name"]}    '
            f'کد پرسنلی: {emp["user_id"]}    نوع عضویت: {emp["membership"]}'
        ),
    )
    cell.font = Font(bold=True, size=11, name=FONT_NAME)
    cell.alignment = Alignment(horizontal='right', vertical='center')
    row += 1
    extra = []
    if emp.get('hire_date_j'):
        extra.append(f'تاریخ استخدام: {emp["hire_date_j"]}')
    if emp.get('termination_date_j'):
        extra.append(f'تاریخ پایان: {emp["termination_date_j"]}')
    if extra:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value='    '.join(extra))
        cell.font = Font(size=10, name=FONT_NAME, color='555555')
        cell.alignment = Alignment(horizontal='right', vertical='center')


def _hourly_leave_and_mission_display(day: dict) -> str:
    """ترکیب مرخصی/ساعتی + مأموریت ساعتی برای ستون «نوع مرخصی» اکسل."""
    parts = []
    leave_name = day.get('leave_name')
    if leave_name:
        parts.append(str(leave_name))
    # Phase 6B: approved hourly mission with start/end/minutes (display only)
    missions = day.get('hourly_missions') or []
    if missions:
        mission_lines = [
            f"مأموریت ساعتی {m['start']}-{m['end']} ({m['minutes']} دقیقه)"
            for m in missions
        ]
        parts.extend(mission_lines)
    elif day.get('hourly_mission_display'):
        parts.append(day['hourly_mission_display'])
    if not parts:
        return '-'
    return '\n'.join(parts)


def _write_daily_table(ws, days: List[dict]):
    header_row = 6
    headers = [
        'تاریخ', 'روز', 'وضعیت روز',
        'وضعیت فرد', 'نوع مرخصی / مرخصی ساعتی', 'ترددها (ورود → خروج)',
    ]
    for column, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=column, value=header)
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
        cell.fill = HEADER_FILL
    ws.row_dimensions[header_row].height = 30

    for row_index, day in enumerate(days, 1):
        row = header_row + row_index
        fill = OFF_FILL if day['day_status'] == 'تعطیل' else (
            EVEN_FILL if row_index % 2 == 0 else None
        )
        leave_display = _hourly_leave_and_mission_display(day)
        values = [
            day['jalali_date'],
            day['day_name'],
            day['day_status'],
            day['person_status_name'],
            leave_display,
            day['attendance_str'],
        ]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.font = DATA_FONT
            cell.border = THIN
            cell.alignment = Alignment(
                horizontal='right' if column == 6 else 'center',
                vertical='center',
                wrap_text=True,
            )
            if fill:
                cell.fill = fill
        ws.row_dimensions[row].height = 34

    guide_row = header_row + len(days) + 2
    ws.merge_cells(start_row=guide_row, start_column=1, end_row=guide_row, end_column=6)
    guide_cell = ws.cell(row=guide_row, column=1, value="راهنمای گزارش: (M) یعنی تردد دستی؛ نبودن (M) یعنی ثبت توسط دستگاه. در هر روز، زمان‌ها به ترتیب ورود و سپس خروج نمایش داده می‌شوند.")
    guide_cell.font = Font(size=9, name=FONT_NAME)
    guide_cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    guide_cell.border = Border(top=THIN_SIDE)
    ws.row_dimensions[guide_row].height = 24

    ws.freeze_panes = 'A7'
    ws.print_title_rows = '1:6'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def _auto_width(ws, widths: List[float]):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _save_workbook(workbook: Workbook, output) :
    if isinstance(output, BytesIO) or hasattr(output, 'write'):
        workbook.save(output)
        output.seek(0)
    else:
        workbook.save(output)
    return output


def export_individual(report: dict, output) :
    workbook = Workbook()
    ws = workbook.active
    ws.title = f"{report['month_name']} {report['year']}"
    ws.sheet_view.rightToLeft = True
    emp = report['employees'][0] if report['employees'] else None
    if emp:
        _write_title(
            ws,
            f"گزارش خام - {emp['full_name']} ({emp['user_id']})",
            report['month_name'],
            report['year'],
            6,
        )
        _write_employee_header(ws, emp)
        _write_daily_table(ws, emp['days'])
    _auto_width(ws, [12, 10, 10, 14, 18, 36])
    return _save_workbook(workbook, output)


def export_group(report: dict, output) :
    workbook = Workbook()
    ws = workbook.active
    ws.title = 'فهرست کارمندان'
    ws.sheet_view.rightToLeft = True
    _write_title(ws, f"گزارش خام - {report['month_name']} {report['year']}", None, None, 5)
    headers = ['ردیف', 'کد', 'نام و نام خانوادگی', 'نوع عضویت', 'شیت']
    for column, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=column, value=header)
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
        cell.fill = HEADER_FILL
    used = {ws.title}
    sheet_names = []
    for emp in report['employees']:
        sheet_names.append(_safe_sheet_name(str(emp['user_id']), used))
    for index, emp in enumerate(report['employees'], 1):
        row = 3 + index
        sheet_name = sheet_names[index - 1]
        values = [index, emp['user_id'], emp['full_name'], emp['membership'], sheet_name]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.font = DATA_FONT
            cell.border = THIN
            cell.alignment = RIGHT if column in (3, 4) else CENTER
            if index % 2 == 0:
                cell.fill = EVEN_FILL
    _auto_width(ws, [6, 14, 24, 18, 24])
    ws.freeze_panes = 'A4'

    for index, emp in enumerate(report['employees']):
        sheet_name = sheet_names[index]
        ws_employee = workbook.create_sheet(sheet_name)
        ws_employee.sheet_view.rightToLeft = True
        _write_title(
            ws_employee,
            f"گزارش خام - {emp['full_name']} ({emp['user_id']})",
            report['month_name'],
            report['year'],
            6,
        )
        _write_employee_header(ws_employee, emp)
        _write_daily_table(ws_employee, emp['days'])
        _auto_width(ws_employee, [12, 10, 10, 14, 18, 36])

    return _save_workbook(workbook, output)
