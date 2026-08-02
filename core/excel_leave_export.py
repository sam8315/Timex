"""
خروجی اکسل گزارش مرخصی‌ها
"""
from datetime import date
from typing import Dict, List
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import jdatetime


class ExcelLeaveExporter:
    """خروجی اکسل گزارش مرخصی‌ها"""

    def __init__(self, output_dir: str = "exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)

    def export_leave_report(self, reports: List[Dict], year: int, month: int, period: str) -> str:
        """خروجی گزارش مرخصی‌ها به اکسل"""
        wb = Workbook()
        ws = wb.active
        ws.title = "گزارش مرخصی‌ها"

        # عنوان
        ws.append([f"گزارش مرخصی‌ها - {period}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=21)
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.append([f"تعداد کاربران: {len(reports)}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=21)
        ws['A2'].alignment = Alignment(horizontal='center')

        ws.append([])

        # ✅ هدر دو سطری
        # سطر اول: گروه‌ها
        headers_row1 = [
            '', '', '',
            'AL', '', '',
            'SL', '', '',
            'RL', '', '',
            'UL', '', '',
            'CW', '', '',
            'TOT', '', ''
        ]
        ws.append(headers_row1)

        # سطر دوم: زیرستون‌ها
        headers_row2 = [
            '#', 'Code', 'Name',
            'T', 'U', 'B',
            'T', 'U', 'B',
            'T', 'U', 'B',
            'T', 'U', 'B',
            'T', 'U', 'B',
            'T', 'U', 'B'
        ]
        ws.append(headers_row2)

        # استایل هدر
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        # رنگ‌های مختلف برای هر گروه
        group_colors = {
            'AL': "4472C4",  # آبی
            'SL': "ED7D31",  # نارنجی
            'RL': "70AD47",  # سبز
            'UL': "FFC000",  # زرد
            'CW': "5B9BD5",  # آبی روشن
            'TOT': "A5A5A5", # خاکستری
        }

        for row_idx in [4, 5]:
            for cell in ws[row_idx]:
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
                cell.fill = header_fill

        # رنگ‌بندی ستون‌ها
        col_groups = [
            (1, 3, "FFFFFF"),  # #, Code, Name
            (4, 6, "4472C4"),  # AL
            (7, 9, "ED7D31"),  # SL
            (10, 12, "70AD47"), # RL
            (13, 15, "FFC000"), # UL
            (16, 18, "5B9BD5"), # CW
            (19, 21, "A5A5A5"), # TOT
        ]

        for start_col, end_col, color in col_groups:
            for col in range(start_col, end_col + 1):
                cell = ws.cell(row=4, column=col)
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        # داده‌ها
        for i, r in enumerate(reports, 1):
            ws.append([
                i,
                r['user_id'],
                r['full_name'],
                # AL
                r['al_total'], r['al_used'], r['al_balance'],
                # SL
                r['sl_total'], r['sl_used'], r['sl_balance'],
                # RL
                r['rl_total'], r['rl_used'], r['rl_balance'],
                # UL
                r['ul_total'], r['ul_used'], r['ul_balance'],
                # CW
                r['cw_total'], r['cw_used'], r['cw_balance'],
                # TOT
                r['tot_total'], r['tot_used'], r['tot_balance'],
            ])

        # تنظیم عرض ستون‌ها
        self._auto_adjust_column_width(ws)

        # ذخیره
        filename = self.output_dir / f"leave_report_{period.replace(' ', '_')}.xlsx"
        wb.save(filename)
        return str(filename)

    def _auto_adjust_column_width(self, ws, max_width: int = 20):
        """تنظیم خودکار عرض ستون‌ها"""
        for col_idx in range(1, ws.max_column + 1):
            column_letter = get_column_letter(col_idx)
            max_length = 0

            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is not None:
                    try:
                        cell_len = len(str(cell.value))
                        if cell_len > max_length:
                            max_length = cell_len
                    except:
                        pass

            adjusted_width = min(max_length + 2, max_width)
            ws.column_dimensions[column_letter].width = max(adjusted_width, 8)