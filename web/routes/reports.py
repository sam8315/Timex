"""
پنل گزارشات مدیریتی
"""
from fastapi import APIRouter, Request, Depends, Form, Query
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from pathlib import Path
from sqlalchemy.orm import Session, selectinload
from typing import Optional
from urllib.parse import quote
from io import BytesIO
import jdatetime

from web.dependencies import get_db, require_admin
from web.permissions import enforce_permission
from models.user import User
from models.employee import Employee
from models.contract import CONTRACT_TYPES
from core.raw_report import (
    JALALI_MONTHS,
    EMPLOYMENT_TYPE_OPTIONS,
    STATUS_FILTER_OPTIONS,
    build_raw_report,
)
from core.excel_raw_report import (
    export_individual as export_raw_excel_individual,
    export_group as export_raw_excel_group,
)
from core.pdf_raw_report import (
    export_individual as export_raw_pdf_individual,
    export_group as export_raw_pdf_group,
)
from core.detailed_monthly_report_v2 import DetailedMonthlyReportGeneratorV2

router = APIRouter(tags=["Reports"])
templates = Jinja2Templates(directory=str(Path(__file__).parent.parent / "templates"))

# نام ماه‌های شمسی
JALALI_MONTHS = {
    1: 'فروردین', 2: 'اردیبهشت', 3: 'خرداد',
    4: 'تیر', 5: 'مرداد', 6: 'شهریور',
    7: 'مهر', 8: 'آبان', 9: 'آذر',
    10: 'دی', 11: 'بهمن', 12: 'اسفند'
}
# 🆕 لیست دیکشنری برای استفاده در تمپلیت‌ها
JALALI_MONTHS_LIST = [
    {'num': num, 'name': name}
    for num, name in JALALI_MONTHS.items()
]


@router.get("/reports", response_class=HTMLResponse)
async def reports_page(
        request: Request,
        user: User = Depends(require_admin),
        db: Session = Depends(get_db)
):
    """صفحه اصلی گزارشات"""
    enforce_permission(db, user, 'view_reports')
    # سال‌های موجود برای انتخاب
    current_year_j = jdatetime.date.today().year
    available_years = list(range(current_year_j, current_year_j - 6, -1))

    return templates.TemplateResponse(request, "admin/reports.html", {
        "user": user,
        "available_years": available_years,
        "current_year": current_year_j,
        "current_month": jdatetime.date.today().month,
        "is_admin": True,
    })


def _raw_employee_search_data(db: Session) -> list:
    employees = db.query(Employee).filter(
        Employee.is_active.is_(True)
    ).order_by(Employee.first_name, Employee.last_name, Employee.user_id).all()
    return [
        {
            'user_id': employee.user_id,
            'full_name': employee.full_name,
            'national_code': employee.national_code or '',
            'card': employee.user_id,
            'department': employee.department or '-',
        }
        for employee in employees
    ]


def _raw_selected_employee(db: Session, user_id: Optional[str]):
    if not user_id:
        return None
    return db.query(Employee).filter(
        Employee.user_id == user_id,
        Employee.is_active.is_(True),
    ).first()


def _raw_context(
    request: Request,
    user: User,
    db: Session,
    report=None,
    selected_user_id: Optional[str] = None,
    selected_year: Optional[int] = None,
    selected_month: Optional[int] = None,
    selected_employment_type: str = 'all',
    selected_status_filter: str = 'all',
):
    today_j = jdatetime.date.today()
    selected_employee = _raw_selected_employee(db, selected_user_id)
    return {
        'user': user,
        'report': report,
        'employees_data': _raw_employee_search_data(db),
        'available_years': list(range(today_j.year, today_j.year - 6, -1)),
        'jalali_months': JALALI_MONTHS,
        'employment_type_options': EMPLOYMENT_TYPE_OPTIONS,
        'status_filter_options': STATUS_FILTER_OPTIONS,
        'selected_user_id': selected_user_id or '',
        'selected_user_name': selected_employee.full_name if selected_employee else '',
        'selected_year': selected_year or today_j.year,
        'selected_month': selected_month or today_j.month,
        'selected_employment_type': selected_employment_type,
        'selected_status_filter': selected_status_filter,
        'is_admin': True,
    }


def _validate_raw_report_params(
    year: int,
    month: int,
    employment_type: str,
    status_filter: str,
):
    if month < 1 or month > 12:
        raise ValueError('ماه نامعتبر است')
    if employment_type not in CONTRACT_TYPES and employment_type != 'all':
        raise ValueError('نوع عضویت نامعتبر است')
    valid_status_filters = {value for value, _ in STATUS_FILTER_OPTIONS}
    if status_filter not in valid_status_filters:
        raise ValueError('فیلتر وضعیت نامعتبر است')


@router.get("/reports/raw", response_class=HTMLResponse)
async def raw_report_form(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    enforce_permission(db, user, 'view_reports')
    today_j = jdatetime.date.today()
    selected_user_id = request.query_params.get('target_user_id', '')
    try:
        selected_year = int(request.query_params.get('year', today_j.year))
        selected_month = int(request.query_params.get('month', today_j.month))
    except (TypeError, ValueError):
        selected_year, selected_month = today_j.year, today_j.month
    selected_employment_type = request.query_params.get('employment_type', 'all')
    selected_status_filter = request.query_params.get('status_filter', 'all')
    _validate_raw_report_params(
        selected_year,
        selected_month,
        selected_employment_type,
        selected_status_filter,
    )
    return templates.TemplateResponse(
        request,
        'admin/report_raw.html',
        _raw_context(
            request,
            user,
            db,
            selected_user_id=selected_user_id,
            selected_year=selected_year,
            selected_month=selected_month,
            selected_employment_type=selected_employment_type,
            selected_status_filter=selected_status_filter,
        ),
    )


@router.post("/reports/raw", response_class=HTMLResponse)
async def raw_report_generate(
    request: Request,
    target_user_id: Optional[str] = Form(None),
    year: int = Form(...),
    month: int = Form(...),
    employment_type: str = Form('all'),
    status_filter: str = Form('all'),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    enforce_permission(db, user, 'view_reports')
    try:
        _validate_raw_report_params(year, month, employment_type, status_filter)
        report = build_raw_report(
            db,
            year,
            month,
            employee_user_id=target_user_id or None,
            employment_type=employment_type,
            status_filter=status_filter,
        )
        return templates.TemplateResponse(
            request,
            'admin/report_raw.html',
            _raw_context(
                request,
                user,
                db,
                report=report,
                selected_user_id=target_user_id,
                selected_year=year,
                selected_month=month,
                selected_employment_type=employment_type,
                selected_status_filter=status_filter,
            ),
        )
    except Exception as e:
        error = quote(str(e), safe='')
        return RedirectResponse(
            url=f'/reports/raw?error={error}',
            status_code=302,
        )


def _raw_export_report(
    db: Session,
    year: int,
    month: int,
    target_user_id: Optional[str],
    employment_type: str,
    status_filter: str,
):
    _validate_raw_report_params(year, month, employment_type, status_filter)
    return build_raw_report(
        db,
        year,
        month,
        employee_user_id=target_user_id or None,
        employment_type=employment_type,
        status_filter=status_filter,
    )


@router.get("/reports/raw/excel")
async def raw_report_excel(
    request: Request,
    year: int = Query(...),
    month: int = Query(...),
    target_user_id: Optional[str] = Query(None),
    employment_type: str = Query('all'),
    status_filter: str = Query('all'),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    enforce_permission(db, user, 'view_reports')
    try:
        report = _raw_export_report(
            db, year, month, target_user_id, employment_type, status_filter
        )
        output = BytesIO()
        if report['mode'] == 'group':
            export_raw_excel_group(report, output)
        else:
            export_raw_excel_individual(report, output)
        filename = f'raw_attendance_{year}_{month:02d}.xlsx'
        utf8_filename = quote(
            f'گزارش_خام_تردد_{year}_{month:02d}.xlsx',
            safe='',
        )
        return StreamingResponse(
            output,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={
                'Content-Disposition': (
                    f'attachment; filename="{filename}"; filename*=UTF-8\'\'{utf8_filename}'
                )
            },
        )
    except Exception as e:
        error = quote(str(e), safe='')
        return RedirectResponse(
            url=f'/reports/raw?error={error}',
            status_code=302,
        )


@router.get("/reports/raw/pdf")
async def raw_report_pdf(
    request: Request,
    year: int = Query(...),
    month: int = Query(...),
    target_user_id: Optional[str] = Query(None),
    employment_type: str = Query('all'),
    status_filter: str = Query('all'),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    enforce_permission(db, user, 'view_reports')
    try:
        report = _raw_export_report(
            db, year, month, target_user_id, employment_type, status_filter
        )
        output = BytesIO()
        if report['mode'] == 'group':
            export_raw_pdf_group(report, output)
        else:
            export_raw_pdf_individual(report, output)
        filename = f'raw_attendance_{year}_{month:02d}.pdf'
        utf8_filename = quote(
            f'گزارش_خام_تردد_{year}_{month:02d}.pdf',
            safe='',
        )
        return StreamingResponse(
            output,
            media_type='application/pdf',
            headers={
                'Content-Disposition': (
                    f'attachment; filename="{filename}"; filename*=UTF-8\'\'{utf8_filename}'
                )
            },
        )
    except Exception as e:
        error = quote(str(e), safe='')
        return RedirectResponse(
            url=f'/reports/raw?error={error}',
            status_code=302,
        )



@router.get("/reports/monthly-detailed", response_class=HTMLResponse)
async def monthly_detailed_report_form(
        request: Request,
        user: User = Depends(require_admin),
        db: Session = Depends(get_db)
):
    """فرم انتخاب کاربر و ماه برای گزارش تفصیلی"""
    enforce_permission(db, user, 'view_reports')
    # دریافت لیست کارمندان فعال
    employees = db.query(Employee).filter(
        Employee.is_active == True
    ).order_by(Employee.first_name, Employee.last_name).all()

    employees_list = [
        {
            'user_id': emp.user_id,
            'full_name': emp.full_name,
            'department': emp.department or '-',
        }
        for emp in employees
    ]

    # سال‌های موجود
    current_year_j = jdatetime.date.today().year
    available_years = list(range(current_year_j, current_year_j - 6, -1))

    return templates.TemplateResponse(request, "admin/report_monthly_detailed.html", {
        "user": user,
        "employees": employees_list,
        "available_years": available_years,
        "jalali_months": JALALI_MONTHS,
        "current_year": current_year_j,
        "current_month": jdatetime.date.today().month,
        "report": None,
        "is_admin": True,
    })


@router.post("/reports/monthly-detailed")
async def monthly_detailed_report_generate(
        request: Request,
        target_user_id: str = Form(...),
        year: int = Form(...),
        month: int = Form(...),
        user: User = Depends(require_admin),
        db: Session = Depends(get_db)
):
    """تولید گزارش تفصیلی ماهانه"""
    enforce_permission(db, user, 'view_reports')
    try:
        # اعتبارسنجی ماه
        if month < 1 or month > 12:
            raise ValueError("ماه نامعتبر است")

        # تولید گزارش با کلاس گزارش‌ساز
        generator = DetailedMonthlyReportGeneratorV2()
        try:
            report = generator.generate_detailed_report(target_user_id, year, month)
        finally:
            generator.close()

        if not report.get('success'):
            raise ValueError(report.get('message', 'خطا در تولید گزارش'))

        # دریافت لیست کارمندان برای فرم
        employees = db.query(Employee).filter(
            Employee.is_active == True
        ).order_by(Employee.first_name, Employee.last_name).all()

        employees_list = [
            {
                'user_id': emp.user_id,
                'full_name': emp.full_name,
                'department': emp.department or '-',
            }
            for emp in employees
        ]

        current_year_j = jdatetime.date.today().year
        available_years = list(range(current_year_j, current_year_j - 6, -1))

        return templates.TemplateResponse(request, "admin/report_monthly_detailed.html", {
            "user": user,
            "employees": employees_list,
            "available_years": available_years,
            "jalali_months": JALALI_MONTHS,
            "current_year": current_year_j,
            "current_month": jdatetime.date.today().month,
            "report": report,
            "selected_user_id": target_user_id,
            "selected_year": year,
            "selected_month": month,
            "is_admin": True,
        })
    except Exception as e:
        return RedirectResponse(
            url=f"/admin/reports/monthly-detailed?error={str(e)}",
            status_code=302
        )


# ============================================
# 🆕 تبدیل ساعت اعشاری به فرمت HH:MM
# ============================================
def format_hhmm(hours):
    """تبدیل ساعت اعشاری به فرمت HH:MM"""
    if hours is None or hours == 0:
        return '-'
    total_minutes = int(round(hours * 60))
    h = total_minutes // 60
    m = total_minutes % 60
    return f"{h:02d}:{m:02d}"


# ثبت filter در Jinja2
templates.env.filters['hhmm'] = format_hhmm


# ============================================
# 🆕 گزارش کامل ماهانه (شبیه PDF)
# ============================================

@router.get("/reports/monthly-full", response_class=HTMLResponse)
async def monthly_full_report_form(
        request: Request,
        user: User = Depends(require_admin),
        db: Session = Depends(get_db)
):
    """فرم گزارش کامل ماهانه (شبیه PDF)"""
    enforce_permission(db, user, 'view_reports')
    employees = db.query(Employee).filter(
        Employee.is_active == True
    ).order_by(Employee.first_name, Employee.last_name).all()

    employees_list = [
        {
            'user_id': emp.user_id,
            'full_name': emp.full_name,
            'department': emp.department or '-',
        }
        for emp in employees
    ]

    current_year_j = jdatetime.date.today().year
    available_years = list(range(current_year_j, current_year_j - 6, -1))

    return templates.TemplateResponse(request, "admin/report_monthly_full.html", {
        "user": user,
        "employees": employees_list,
        "available_years": available_years,
        "jalali_months": JALALI_MONTHS,
        "current_year": current_year_j,
        "current_month": jdatetime.date.today().month,
        "report": None,
        "is_admin": True,
    })


@router.post("/reports/monthly-full")
async def monthly_full_report_generate(
        request: Request,
        target_user_id: str = Form(...),
        year: int = Form(...),
        month: int = Form(...),
        user: User = Depends(require_admin),
        db: Session = Depends(get_db)
):
    """تولید گزارش کامل ماهانه (شبیه PDF)"""
    enforce_permission(db, user, 'view_reports')
    try:
        if month < 1 or month > 12:
            raise ValueError("ماه نامعتبر است")

        generator = DetailedMonthlyReportGeneratorV2()
        try:
            report = generator.generate_detailed_report(target_user_id, year, month)
        finally:
            generator.close()

        if not report.get('success'):
            raise ValueError(report.get('message', 'خطا در تولید گزارش'))

        employees = db.query(Employee).filter(
            Employee.is_active == True
        ).order_by(Employee.first_name, Employee.last_name).all()

        employees_list = [
            {
                'user_id': emp.user_id,
                'full_name': emp.full_name,
                'department': emp.department or '-',
            }
            for emp in employees
        ]

        current_year_j = jdatetime.date.today().year
        available_years = list(range(current_year_j, current_year_j - 6, -1))

        return templates.TemplateResponse(request, "admin/report_monthly_full.html", {
            "user": user,
            "employees": employees_list,
            "available_years": available_years,
            "jalali_months": JALALI_MONTHS,
            "current_year": current_year_j,
            "current_month": jdatetime.date.today().month,
            "report": report,
            "selected_user_id": target_user_id,
            "selected_year": year,
            "selected_month": month,
            "is_admin": True,
        })
    except Exception as e:
        return RedirectResponse(
            url=f"/reports/monthly-full?error={str(e)}",
            status_code=302
        )


# ============================================
# 📊 گزارش آمار ماهیانه
# ============================================
from io import BytesIO
from models.attendance import Attendance
from models.contract import Contract
from models.holiday import Holiday
from models.daily_status import DailyStatus
from models.leave_request import LeaveRequest
from web.services.travel_leave_service import build_leave_days_by_date
from sqlalchemy import and_, func, or_
from datetime import timedelta


DEPT_NAMES_REPORT = {
    '1': 'رسمی', '2': 'وظیفه', '3': 'خریدخدمت',
    '4': 'قراردادی', '5': 'پزشک'
}


def get_month_days_count(year: int, month: int) -> int:
    """تعداد روزهای ماه شمسی (اسفند در سال کبیسه ۳۰ روز)"""
    if month <= 6:
        return 31
    elif month <= 11:
        return 30
    else:
        try:
            return 30 if jdatetime.date(year, 1, 1).isleap() else 29
        except Exception:
            return 29


def get_month_end_jalali(year: int, month: int):
    """آخرین روز ماه شمسی (با لحاظ کبیسه برای اسفند)"""
    if month == 12:
        return jdatetime.date(year, 12, get_month_days_count(year, 12))
    return jdatetime.date(year, month + 1, 1) - timedelta(days=1)


def get_day_code(
        day_date,
        leaves_by_date: dict,
        daily_status_map: dict,
        holiday_dates: set,
        has_attendance: bool,
        is_friday: bool,
        hire_date=None,  # 🆕 تاریخ عضویت
        contract_end_date=None  # 🆕 تاریخ پایان قرارداد
) -> str:
    """تعیین کد یک روز برای گزارش آمار ماهیانه"""

    # 🆕 اولویت ۱: معرفی (تاریخ عضویت)
    if hire_date and day_date == hire_date:
        return 'معرفی'

    # 🆕 اولویت ۲: تسویه (تاریخ پایان قرارداد)
    if contract_end_date and day_date == contract_end_date:
        return 'تسویه'

    # اولویت ۳: جمعه یا تعطیل رسمی → اگر تردد دارد حاضر، وگرنه خالی
    if is_friday or day_date in holiday_dates:
        if has_attendance:
            return '✓'
        return ''

    # اولویت ۴: مرخصی تأیید شده (فقط در روزهای کاری)
    leave_type = leaves_by_date.get(day_date)
    if leave_type:
        if leave_type == 'AL':
            return 'ص'
        elif leave_type == 'SL':
            return 'ج'
        elif leave_type == 'RL':
            return 'ت'
        elif leave_type == 'UL':
            return 'ب'
        elif leave_type == 'CW':
            return 'ذ'
        elif leave_type == 'TL':
            return 'TL'

    # اولویت ۵: وضعیت دستی (DailyStatus)
    status_code = daily_status_map.get(day_date)
    if isinstance(status_code, str):
        status_code = status_code.strip()
    if status_code:
        if status_code == 'A':
            return 'غ'
        elif status_code == 'R':
            return 'اس'
        elif status_code == 'M':
            return 'م'

    # اولویت ۶ و ۷: حضور یا بدون تردد
    if has_attendance:
        return '✓'
    else:
        return '-'


DAY_CODE_DISPLAY = {
    'TL': 'تو',
}


def display_day_code(code) -> str:
    """Persian presentation for internal day codes (keeps backend code intact)."""
    return DAY_CODE_DISPLAY.get(code, code)


templates.env.filters['day_code_display'] = display_day_code


def resolve_intro_settle_dates(user_contracts, hire_date, termination_date,
                               month_start_g, month_end_g):
    """تعیین تاریخ معرفی و تسویه یک کارمند برای بازه ماه.

    - اگر قرارداد مرتبط با ماه پیدا شود: معرفی = شروع قرارداد (در ماه)،
      تسویه = پایان قرارداد (در ماه) مگر اینکه قرارداد جدیدی بلافاصله بعد
      شروع شده باشد.
    - اگر قراردادی پیدا نشود: معرفی = تاریخ استخدام، تسویه = تاریخ ترک کار.
    """
    if user_contracts:
        starts_in_month = [
            c.start_date for c in user_contracts
            if month_start_g <= c.start_date <= month_end_g
        ]
        intro_date = min(starts_in_month) if starts_in_month else None

        ends_in_month = [
            c.end_date for c in user_contracts
            if c.end_date and month_start_g <= c.end_date <= month_end_g
        ]
        settle_date = None
        for end_date in sorted(ends_in_month):
            next_day = end_date + timedelta(days=1)
            continued = any(
                c.start_date <= next_day
                and (c.end_date is None or c.end_date >= next_day)
                for c in user_contracts
            )
            if not continued:
                settle_date = end_date
        return intro_date, settle_date

    return hire_date, termination_date


def fetch_monthly_stats_leave_data(db: Session, month_start_g, month_end_g):
    """Approved day-leaves overlapping the month plus holidays over the full
    leave ranges (needed so Travel Leave counting sees days before the month).
    """
    approved_leaves = db.query(LeaveRequest).options(
        selectinload(LeaveRequest.travel_leave_detail)
    ).filter(
        and_(
            LeaveRequest.status == 'A',
            LeaveRequest.leave_type != 'HL',
            LeaveRequest.from_date <= month_end_g,
            LeaveRequest.to_date >= month_start_g
        )
    ).all()

    holiday_start = month_start_g
    holiday_end = month_end_g
    if approved_leaves:
        holiday_start = min(
            holiday_start,
            min(leave.from_date for leave in approved_leaves)
        )
        holiday_end = max(
            holiday_end,
            max(leave.to_date for leave in approved_leaves)
        )

    holidays = db.query(Holiday).filter(
        and_(
            Holiday.holiday_date >= holiday_start,
            Holiday.holiday_date <= holiday_end
        )
    ).all()
    return approved_leaves, holidays


def display_holiday_dates(holidays, month_start_g, month_end_g) -> set:
    """Holiday dates shown in the report: month window only, all groups."""
    return {
        h.holiday_date for h in holidays
        if month_start_g <= h.holiday_date <= month_end_g
    }


def build_monthly_stats_leaves_map(
    approved_leaves,
    holidays,
    departments_by_user,
    month_start_g,
    month_end_g,
):
    """``{user_id: {date: leave_type}}`` with the Travel Leave split applied.

    Each user is processed independently through ``build_leave_days_by_date``
    so leave dates of different employees never interfere. Holidays follow the
    existing group semantics: national (``group_id is None``) or matching the
    employee's department.
    """
    leaves_by_user = {}
    for leave in approved_leaves:
        leaves_by_user.setdefault(leave.user_id, []).append(leave)

    leaves_map = {}
    holiday_cache = {}
    for uid, user_leaves in leaves_by_user.items():
        department = departments_by_user.get(uid)
        if department not in holiday_cache:
            holiday_cache[department] = {
                h.holiday_date for h in holidays
                if h.group_id is None or h.group_id == department
            }
        leaves_map[uid] = build_leave_days_by_date(
            user_leaves,
            holiday_cache[department],
            month_start_g,
            month_end_g,
        )
    return leaves_map


@router.get("/reports/monthly-stats", response_class=HTMLResponse)
async def monthly_stats_report_form(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    """فرم گزارش آمار ماهیانه"""
    enforce_permission(db, user, 'view_reports')
    current_year_j = jdatetime.date.today().year
    available_years = list(range(current_year_j, current_year_j - 6, -1))

    return templates.TemplateResponse(request, "admin/report_monthly_stats.html", {
        "user": user,
        "available_years": available_years,
        "jalali_months": JALALI_MONTHS_LIST,  # ✅ لیست دیکشنری
        "current_year": current_year_j,
        "current_month": jdatetime.date.today().month,
        "dept_names": DEPT_NAMES_REPORT,
        "report": None,
        "selected_department": "all",
        "is_admin": True,
    })


@router.post("/reports/monthly-stats")
async def monthly_stats_report_generate(
    request: Request,
    year: int = Form(...),
    month: int = Form(...),
    department: str = Form("all"),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    """تولید گزارش آمار ماهیانه"""
    enforce_permission(db, user, 'view_reports')
    try:
        if month < 1 or month > 12:
            raise ValueError("ماه نامعتبر است")

        # بازه ماه شمسی → میلادی (با لحاظ کبیسه اسفند)
        month_start_j = jdatetime.date(year, month, 1)
        month_end_j = get_month_end_jalali(year, month)
        month_start_g = month_start_j.togregorian()
        month_end_g = month_end_j.togregorian()
        days_count = get_month_days_count(year, month)

        # دریافت کارمندان فعال در بازه ماه
        emp_query = db.query(Employee).filter(
            Employee.hire_date <= month_end_g,
            or_(
                Employee.termination_date == None,
                Employee.termination_date >= month_start_g
            )
        )
        if department != "all":
            emp_query = emp_query.filter(Employee.department == department)

        # مرتب‌سازی بر اساس تاریخ عضویت
        employees = emp_query.order_by(Employee.hire_date, Employee.first_name).all()

        # مرخصی‌های تأیید شده + تعطیلات بازه کامل مرخصی‌ها (برای جداسازی TL)
        approved_leaves, all_holidays = fetch_monthly_stats_leave_data(
            db, month_start_g, month_end_g
        )

        # ساخت دیکشنری مرخصی‌ها بر اساس کاربر و تاریخ (هر کاربر مستقل)
        leaves_map = build_monthly_stats_leaves_map(
            approved_leaves,
            all_holidays,
            {emp.user_id: emp.department for emp in employees},
            month_start_g,
            month_end_g,
        )

        # دریافت وضعیت‌های دستی (DailyStatus)
        daily_statuses = db.query(DailyStatus).filter(
            and_(
                DailyStatus.status_date >= month_start_g,
                DailyStatus.status_date <= month_end_g
            )
        ).all()
        status_map = {}  # {user_id: {date: status_code}}
        for ds in daily_statuses:
            if ds.user_id not in status_map:
                status_map[ds.user_id] = {}
            status_map[ds.user_id][ds.status_date] = ds.status_code

        # دریافت ترددها در بازه ماه
        attendances = db.query(Attendance).filter(
            and_(
                Attendance.timestamp >= month_start_g - timedelta(days=1),
                Attendance.timestamp <= month_end_g + timedelta(days=1),
                Attendance.is_deleted == False
            )
        ).all()
        att_map = {}  # {user_id: {date: True}}
        for att in attendances:
            uid = att.user_id
            att_date = att.timestamp.date()
            if uid not in att_map:
                att_map[uid] = set()
            att_map[uid].add(att_date)

        # دریافت تعطیلات رسمی (فقط بازه ماه برای نمایش)
        holiday_dates = display_holiday_dates(
            all_holidays, month_start_g, month_end_g
        )

        # دریافت قراردادهای مرتبط با ماه (یک کوئری برای همه کاربران؛
        # تا فردای پایان ماه تا قرارداد بعدیِ بلافاصله بعد هم دیده شود)
        month_contracts = db.query(Contract).filter(
            and_(
                Contract.start_date <= month_end_g + timedelta(days=1),
                or_(Contract.end_date == None,
                    Contract.end_date >= month_start_g)
            )
        ).all()
        contracts_by_user = {}
        for c in month_contracts:
            contracts_by_user.setdefault(c.user_id, []).append(c)

        # ساخت داده‌های گزارش
        report_rows = []
        for idx, emp in enumerate(employees, 1):
            uid = emp.user_id
            user_leaves = leaves_map.get(uid, {})
            user_statuses = status_map.get(uid, {})
            user_att = att_map.get(uid, set())

            # 🆕 تاریخ معرفی و تسویه: از قرارداد، وگرنه از استخدام/ترک کار
            emp_hire_date, emp_contract_end = resolve_intro_settle_dates(
                contracts_by_user.get(uid, []),
                emp.hire_date if emp.hire_date else None,
                emp.termination_date if emp.termination_date else None,
                month_start_g, month_end_g
            )

            # ستون‌های روزها
            day_codes = []
            for day_num in range(1, days_count + 1):
                day_date = month_start_g + timedelta(days=day_num - 1)
                is_friday = day_date.weekday() == 4

                code = get_day_code(
                    day_date=day_date,
                    leaves_by_date=user_leaves,
                    daily_status_map=user_statuses,
                    holiday_dates=holiday_dates,
                    has_attendance=(day_date in user_att),
                    is_friday=is_friday,
                    hire_date=emp_hire_date,  # 🆕
                    contract_end_date=emp_contract_end  # 🆕
                )
                day_codes.append(code)

            dept_name = DEPT_NAMES_REPORT.get(emp.department, emp.department or '-')

            report_rows.append({
                'row_num': idx,
                'first_name': emp.first_name or '',
                'last_name': emp.last_name or '',
                'department': dept_name,
                'day_codes': day_codes,
                # 🆕 last_col حذف شد
            })
        current_year_j = jdatetime.date.today().year
        available_years = list(range(current_year_j, current_year_j - 6, -1))

        # 🆕 محاسبه روزهای جمعه و تعطیل برای بک‌گراند قرمز
        friday_or_holiday_days = []
        for d in range(1, days_count + 1):
            day_date = month_start_g + timedelta(days=d - 1)
            is_friday = day_date.weekday() == 4
            is_holiday = day_date in holiday_dates
            if is_friday or is_holiday:
                friday_or_holiday_days.append(d)


        return templates.TemplateResponse(request, "admin/report_monthly_stats.html", {
            "user": user,
            "available_years": available_years,
            "jalali_months": JALALI_MONTHS_LIST,  # ✅ لیست دیکشنری (نه JALALI_MONTHS)
            "current_year": year,
            "current_month": month,
            "dept_names": DEPT_NAMES_REPORT,
            "report": {
                "year": year,
                "month": month,
                "month_name": JALALI_MONTHS.get(month, ""),
                "days_count": days_count,
                "rows": report_rows,
                "total_employees": len(report_rows),
            },
            "selected_department": department,
            "is_admin": True,
            "friday_or_holiday_days": friday_or_holiday_days,
        })

    except Exception as e:
        return RedirectResponse(
            url=f"/reports/monthly-stats?error={str(e)}",
            status_code=302
        )


@router.get("/reports/monthly-stats/excel")
async def monthly_stats_report_excel(
    request: Request,
    year: int = Query(...),
    month: int = Query(...),
    department: str = Query("all"),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db)
):
    """خروجی اکسل گزارش آمار ماهیانه - با رنگ‌بندی کامل"""
    enforce_permission(db, user, 'view_reports')
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, Alignment, Border, Side, PatternFill, NamedStyle
        )
        from fastapi.responses import StreamingResponse

        # ============================================
        # 📅 محاسبه بازه ماه (با لحاظ کبیسه اسفند)
        # ============================================
        month_start_j = jdatetime.date(year, month, 1)
        month_end_j = get_month_end_jalali(year, month)
        month_start_g = month_start_j.togregorian()
        month_end_g = month_end_j.togregorian()
        days_count = get_month_days_count(year, month)

        # ============================================
        # 📋 دریافت کارمندان
        # ============================================
        emp_query = db.query(Employee).filter(
            Employee.hire_date <= month_end_g,
            or_(
                Employee.termination_date == None,
                Employee.termination_date >= month_start_g
            )
        )
        if department != "all":
            emp_query = emp_query.filter(Employee.department == department)
        employees = emp_query.order_by(
            Employee.hire_date, Employee.first_name
        ).all()

        # ============================================
        # 🗂️ دریافت داده‌های مورد نیاز
        # ============================================
        # مرخصی‌های تایید شده + تعطیلات بازه کامل مرخصی‌ها (برای جداسازی TL)
        approved_leaves, all_holidays = fetch_monthly_stats_leave_data(
            db, month_start_g, month_end_g
        )
        leaves_map = build_monthly_stats_leaves_map(
            approved_leaves,
            all_holidays,
            {emp.user_id: emp.department for emp in employees},
            month_start_g,
            month_end_g,
        )

        # وضعیت‌های دستی
        daily_statuses = db.query(DailyStatus).filter(
            and_(
                DailyStatus.status_date >= month_start_g,
                DailyStatus.status_date <= month_end_g
            )
        ).all()
        status_map = {}
        for ds in daily_statuses:
            if ds.user_id not in status_map:
                status_map[ds.user_id] = {}
            status_map[ds.user_id][ds.status_date] = ds.status_code

        # ترددها
        attendances = db.query(Attendance).filter(
            and_(
                Attendance.timestamp >= month_start_g - timedelta(days=1),
                Attendance.timestamp <= month_end_g + timedelta(days=1),
                Attendance.is_deleted == False
            )
        ).all()
        att_map = {}
        for att in attendances:
            uid = att.user_id
            att_date = att.timestamp.date()
            if uid not in att_map:
                att_map[uid] = set()
            att_map[uid].add(att_date)

        # تعطیلات (فقط بازه ماه برای نمایش)
        holiday_dates = display_holiday_dates(
            all_holidays, month_start_g, month_end_g
        )

        # قراردادهای مرتبط با ماه (یک کوئری برای همه کاربران)
        month_contracts = db.query(Contract).filter(
            and_(
                Contract.start_date <= month_end_g + timedelta(days=1),
                or_(Contract.end_date == None,
                    Contract.end_date >= month_start_g)
            )
        ).all()
        contracts_by_user = {}
        for c in month_contracts:
            contracts_by_user.setdefault(c.user_id, []).append(c)

        # ============================================
        # 🎨 تعریف استایل‌های رنگی
        # ============================================
        # فونت‌ها
        FONT_NAME = 'Tahoma'
        header_font = Font(bold=True, size=10, name=FONT_NAME, color='000000')
        header_holiday_font = Font(bold=True, size=10, name=FONT_NAME, color='DC3545')
        data_font = Font(size=10, name=FONT_NAME)

        # فونت‌های رنگی برای کدها
        font_present = Font(size=11, name=FONT_NAME, bold=True, color='0D6E0D')      # ✓ سبز پررنگ
        font_absent = Font(size=10, name=FONT_NAME, color='DC3545')                  # - قرمز
        font_leave_al = Font(size=10, name=FONT_NAME, bold=True, color='0D6EFD')    # ص آبی
        font_leave_sl = Font(size=10, name=FONT_NAME, bold=True, color='6F42C1')    # ج بنفش
        font_leave_rl = Font(size=10, name=FONT_NAME, bold=True, color='FD7E14')    # ت نارنجی
        font_leave_travel = Font(size=10, name=FONT_NAME, bold=True, color='D63384')
        font_rest = Font(size=10, name=FONT_NAME, color='6C757D')                   # اس خاکستری
        font_mission = Font(size=10, name=FONT_NAME, bold=True, color='20C997')     # م سبزآبی
        font_gheyb = Font(size=10, name=FONT_NAME, bold=True, color='DC3545')       # غ قرمز پررنگ
        font_moarefa = Font(size=9, name=FONT_NAME, bold=True, color='0DCAF0')      # معرفی فیروزه‌ای
        font_tasvieh = Font(size=9, name=FONT_NAME, bold=True, color='6610F2')      # تسویه بنفش تیره

        # Alignment
        center_align = Alignment(
            horizontal='center', vertical='center', wrap_text=False
        )

        # Border نازک
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )

        # PatternFill (بک‌گراند)
        fill_header = PatternFill(
            start_color='D9E1F2', end_color='D9E1F2', fill_type='solid'
        )  # خاکستری-آبی روشن برای هدر
        fill_holiday_header = PatternFill(
            start_color='F8D7DA', end_color='F8D7DA', fill_type='solid'
        )  # قرمز کم‌رنگ برای سرستون جمعه/تعطیل
        fill_holiday_cell = PatternFill(
            start_color='FFF5F5', end_color='FFF5F5', fill_type='solid'
        )  # قرمز خیلی کم‌رنگ برای سلول جمعه/تعطیل
        fill_moarefa = PatternFill(
            start_color='E0F7FA', end_color='E0F7FA', fill_type='solid'
        )  # فیروزه‌ای کم‌رنگ برای معرفی
        fill_tasvieh = PatternFill(
            start_color='EDE7F6', end_color='EDE7F6', fill_type='solid'
        )  # بنفش کم‌رنگ برای تسویه
        fill_even_row = PatternFill(
            start_color='FAFAFA', end_color='FAFAFA', fill_type='solid'
        )  # خاکستری خیلی روشن برای ردیف‌های زوج

        # ============================================
        # 📋 تعیین روزهای جمعه و تعطیل
        # ============================================
        holiday_day_numbers = set()
        for d in range(1, days_count + 1):
            day_date = month_start_g + timedelta(days=d - 1)
            if day_date.weekday() == 4 or day_date in holiday_dates:
                holiday_day_numbers.add(d)

        # ============================================
        # 📝 ساخت فایل اکسل
        # ============================================
        wb = openpyxl.Workbook()
        ws = wb.active
        month_name = JALALI_MONTHS.get(month, str(month))
        ws.title = f"{year}-{month:02d}-{month_name}"

        # تنظیم جهت صفحه (RTL)
        ws.sheet_view.rightToLeft = True

        # ============================================
        # 📑 ردیف عنوان گزارش (ردیف ۱)
        # ============================================
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days_count + 4)
        title_cell = ws.cell(row=1, column=1)
        title_cell.value = f"گزارش آمار ماهیانه - {month_name} {year}"
        title_cell.font = Font(bold=True, size=14, name=FONT_NAME, color='000000')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        title_cell.fill = PatternFill(
            start_color='4472C4', end_color='4472C4', fill_type='solid'
        )
        title_cell.font = Font(bold=True, size=14, name=FONT_NAME, color='FFFFFF')
        ws.row_dimensions[1].height = 30

        # ============================================
        # 📑 ردیف هدر ستون‌ها (ردیف ۲)
        # ============================================
        headers = ['ردیف', 'نام', 'نام خانوادگی', 'نوع عضویت']
        for d in range(1, days_count + 1):
            headers.append(str(d))

        header_row = 2
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

            # رنگ هدر روزهای جمعه/تعطیل
            if col_idx > 4 and (col_idx - 4) in holiday_day_numbers:
                cell.fill = fill_holiday_header
                cell.font = header_holiday_font
            else:
                cell.fill = fill_header

        ws.row_dimensions[header_row].height = 20

        # ============================================
        # 📝 ردیف‌های داده (از ردیف ۳)
        # ============================================
        for idx, emp in enumerate(employees, 1):
            uid = emp.user_id
            user_leaves = leaves_map.get(uid, {})
            user_statuses = status_map.get(uid, {})
            user_att = att_map.get(uid, set())

            # 🆕 تاریخ معرفی و تسویه: از قرارداد، وگرنه از استخدام/ترک کار
            emp_hire_date, emp_contract_end = resolve_intro_settle_dates(
                contracts_by_user.get(uid, []),
                emp.hire_date if emp.hire_date else None,
                emp.termination_date if emp.termination_date else None,
                month_start_g, month_end_g
            )

            row_num = idx + 2  # شروع از ردیف 3
            is_even_row = (idx % 2 == 0)

            # ستون ردیف
            cell = ws.cell(row=row_num, column=1, value=idx)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border
            if is_even_row:
                cell.fill = fill_even_row

            # ستون نام
            cell = ws.cell(row=row_num, column=2, value=emp.first_name or '')
            cell.font = data_font
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.border = thin_border
            if is_even_row:
                cell.fill = fill_even_row

            # ستون نام خانوادگی
            cell = ws.cell(row=row_num, column=3, value=emp.last_name or '')
            cell.font = data_font
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.border = thin_border
            if is_even_row:
                cell.fill = fill_even_row

            # ستون نوع عضویت
            dept_name = DEPT_NAMES_REPORT.get(
                emp.department, emp.department or '-'
            )
            cell = ws.cell(row=row_num, column=4, value=dept_name)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border
            if is_even_row:
                cell.fill = fill_even_row

            # ستون‌های روزها
            for day_num in range(1, days_count + 1):
                day_date = month_start_g + timedelta(days=day_num - 1)
                is_friday = day_date.weekday() == 4
                is_holiday_day = day_date in holiday_dates

                code = get_day_code(
                    day_date, user_leaves, user_statuses,
                    holiday_dates, (day_date in user_att), is_friday,
                    hire_date=emp_hire_date,
                    contract_end_date=emp_contract_end
                )

                col_idx = day_num + 4  # بعد از 4 ستون ثابت
                cell = ws.cell(
                    row=row_num, column=col_idx,
                    value=display_day_code(code)
                )
                cell.alignment = center_align
                cell.border = thin_border

                # 🎨 بک‌گراند قرمز برای جمعه/تعطیل
                if day_num in holiday_day_numbers:
                    cell.fill = fill_holiday_cell

                # 🎨 رنگ و فونت بر اساس کد
                if code == '✓':
                    cell.font = font_present
                elif code == '-':
                    cell.font = font_absent
                elif code == 'ص':
                    cell.font = font_leave_al
                elif code == 'ج':
                    cell.font = font_leave_sl
                elif code == 'ت':
                    cell.font = font_leave_rl
                elif code == 'TL':
                    cell.font = font_leave_travel
                elif code == 'اس':
                    cell.font = font_rest
                elif code == 'م':
                    cell.font = font_mission
                elif code == 'غ':
                    cell.font = font_gheyb
                elif code == 'معرفی':
                    cell.font = font_moarefa
                    cell.fill = fill_moarefa  # بک‌گراند فیروزه‌ای
                elif code == 'تسویه':
                    cell.font = font_tasvieh
                    cell.fill = fill_tasvieh  # بک‌گراند بنفش
                else:
                    cell.font = data_font

        # ============================================
        # 📐 تنظیم عرض ستون‌ها
        # ============================================
        ws.column_dimensions['A'].width = 6    # ردیف
        ws.column_dimensions['B'].width = 14   # نام
        ws.column_dimensions['C'].width = 16   # نام خانوادگی
        ws.column_dimensions['D'].width = 12   # نوع عضویت
        for d in range(1, days_count + 1):
            col_letter = openpyxl.utils.get_column_letter(d + 4)
            ws.column_dimensions[col_letter].width = 5  # روزها

        # ============================================
        # 📌 ردیف راهنما (آخرین ردیف)
        # ============================================
        last_row = len(employees) + 4
        ws.merge_cells(start_row=last_row, start_column=1,
                      end_row=last_row, end_column=days_count + 4)
        guide_cell = ws.cell(row=last_row, column=1)
        guide_cell.value = (
            "راهنما: ✓=حاضر | -=بدون تردد | ص=استحقاقی | ج=استعلاجی | "
            "ت=تشویقی | تو=توراهی | غ=غایب | اس=استراحت | م=مأموریت | "
            "معرفی=شروع قرارداد (یا عضویت) | تسویه=پایان قرارداد (یا ترک کار) | خالی=جمعه/تعطیل"
        )
        guide_cell.font = Font(size=9, name=FONT_NAME, italic=True, color='666666')
        guide_cell.alignment = Alignment(horizontal='right', vertical='center')
        guide_cell.fill = PatternFill(
            start_color='F8F9FA', end_color='F8F9FA', fill_type='solid'
        )

        # ============================================
        # 📌 ثابت نگه داشتن سرستون‌ها هنگام اسکرول
        # ============================================
        ws.freeze_panes = 'E3'  # از ستون E (اولین روز) و ردیف 3 ثابت

        # ============================================
        # 💾 ذخیره در حافظه و ارسال
        # ============================================
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # 🆕 نام فایل با کاراکترهای فارسی → URL encode
        from urllib.parse import quote
        month_name_en = {
            1: 'Farvardin', 2: 'Ordibehesht', 3: 'Khordad',
            4: 'Tir', 5: 'Mordad', 6: 'Shahrivar',
            7: 'Mehr', 8: 'Aban', 9: 'Azar',
            10: 'Dey', 11: 'Bahman', 12: 'Esfand'
        }.get(month, str(month))

        # نام فایل ساده انگلیسی (برای سازگاری با همه مرورگرها)
        ascii_filename = f"attendance_stats_{year}_{month:02d}_{month_name_en}.xlsx"
        # نام فایل UTF-8 با encoding صحیح (RFC 5987)
        utf8_filename = quote(f"گزارش_آمار_{year}_{month:02d}_{month_name}.xlsx")

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=\"{ascii_filename}\"; filename*=UTF-8''{utf8_filename}"
            }
        )

    except ImportError:
        return RedirectResponse(
            url="/reports/monthly-stats?error=کتابخانه openpyxl نصب نیست",
            status_code=302
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        return RedirectResponse(
            url=f"/reports/monthly-stats?error={str(e)}",
            status_code=302
        )